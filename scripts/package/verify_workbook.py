#!/usr/bin/env python3
"""Re-opens the analysis workbook of an EXTRACTED package and checks it against the package's own CSV files and the stored bundle.

  python -X utf8 verify_workbook.py --package-dir DIR --bundle bundle.json --scope KR|JP|ALL --out report.json

Checks (all must pass, exit code 0): expected sheets and no extras; S/P/D/δ/legacy matrices 47×47 equal to the CSV matrices and to the
bundle's S (exact floats, not rounded); the cached results of the formulas on the check sheet (and that nothing else is a formula);
no error values; ids and ZIP-style text stay text; coordinates, cluster membership, pair diagnostics, dimension and k tables equal their
CSV files; the response sheet holds exactly the scope's participants (KR file: only KR; JP file: only JP; ALL: both) with 47 rows each;
statement sheet has only the scope's language columns; no personal-data columns; charts are native and have no external link.
"""
import argparse
import csv
import glob
import json
import os
import re
import sys
import zipfile

import numpy as np
import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("--package-dir", required=True)
ap.add_argument("--bundle", required=True)
ap.add_argument("--scope", required=True)
ap.add_argument("--out", required=True)
a = ap.parse_args()

root = glob.glob(os.path.join(a.package_dir, "*"))[0] if not os.path.exists(os.path.join(a.package_dir, "manifest.json")) else a.package_dir
bundle = json.load(open(a.bundle, encoding="utf-8"))
bundle = bundle.get("bundle", bundle)
S_bundle = np.array(bundle["analysis"]["S"], dtype=float)
N = bundle["analysis"]["N"]
n = len(bundle["statements"])
xlsx = glob.glob(os.path.join(root, "01_요약", "*_분석자료.xlsx"))
results = []


def check(name, ok, detail=""):
    results.append({"check": name, "pass": bool(ok), "detail": str(detail)[:300]})


check("exactly one analysis workbook", len(xlsx) == 1, xlsx)
path = xlsx[0]
wb_f = openpyxl.load_workbook(path, data_only=False)
wb_v = openpyxl.load_workbook(path, data_only=True)
names = wb_f.sheetnames
check("sheet names unique and <= 31 chars", len(set(names)) == len(names) and all(len(x) <= 31 for x in names), names)
required = ["안내·데이터사전", "메타데이터", "참여자_분류응답(가명)", "참여자_묶음요약", "진술문_원문대응", "S_유사성빈도", "P_유사성비율", "D_비유사성빈도", "delta_정규화비유사성", "기존_빈도행렬_대각0(유사성)", "행렬정의·수식검산", "좌표·군집(추천)", "추천군집·문항", "군집소속_k별", "차원별_지표", "군집수별_지표", "Ward병합이력", "문항쌍_진단", "문항별_stress기여", "안정성_분할반분", "분할반분_전체값", "포함제외·검증", "지표정의·해석범위", "참고문헌", "그림목록", "차트"]
if a.scope == "ALL":
    required += ["한일비교", "S_한국", "S_일본", "P_한국", "P_일본", "P차이_한국-일본"]
check("all required sheets present", all(x in names for x in required), [x for x in required if x not in names])
check("no unexpected extra sheets", set(names) <= set(required), sorted(set(names) - set(required)))


def read_csv(rel):
    with open(os.path.join(root, rel), encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def sheet_rows(ws, header_row=1):
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    return rows


def matrix(ws):
    vals = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=n + 1, min_col=2, max_col=n + 1)]
    return np.array(vals, dtype=float)


def mat_labels_ok(ws):
    top = [ws.cell(1, j + 2).value for j in range(n)]
    left = [ws.cell(i + 2, 1).value for i in range(n)]
    want = [f"S{i + 1:02d}" for i in range(n)]
    return top == want and left == want


def csv_matrix(rel):
    rows = read_csv(rel)
    return np.array([[float(x) for x in r[1:]] for r in rows[1:]], dtype=float)


mats = {"S_유사성빈도": "04_행렬/S_집단유사성빈도_대각선N.csv", "P_유사성비율": "04_행렬/P_유사성비율_대각선1.csv", "D_비유사성빈도": "04_행렬/D_비유사성빈도_대각선0.csv", "delta_정규화비유사성": "04_행렬/delta_비유사성비율_MDS입력_대각선0.csv", "기존_빈도행렬_대각0(유사성)": "04_행렬/기존표시형식_빈도행렬_대각선0.csv"}
M = {}
for sh, rel in mats.items():
    ws = wb_v[sh]
    M[sh] = matrix(ws)
    check(f"{sh}: 47x47, labels S01..S47 in order", M[sh].shape == (n, n) and mat_labels_ok(ws))
    check(f"{sh}: equals CSV {os.path.basename(rel)} exactly", np.array_equal(M[sh], csv_matrix(rel)))
check("S equals the bundle's S", np.array_equal(M["S_유사성빈도"], S_bundle))
check("S symmetric, diagonal = N, 0<=S<=N", np.array_equal(M["S_유사성빈도"], M["S_유사성빈도"].T) and np.all(np.diag(M["S_유사성빈도"]) == N) and M["S_유사성빈도"].min() >= 0 and M["S_유사성빈도"].max() <= N)
check("P = S/N (diagonal 1)", np.array_equal(M["P_유사성비율"], M["S_유사성빈도"] / N) and np.all(np.diag(M["P_유사성비율"]) == 1))
check("D = N - S (diagonal 0)", np.array_equal(M["D_비유사성빈도"], N - M["S_유사성빈도"]) and np.all(np.diag(M["D_비유사성빈도"]) == 0))
dl = 1 - M["S_유사성빈도"] / N
np.fill_diagonal(dl, 0)
check("delta = 1 - S/N (diagonal 0)", np.array_equal(M["delta_정규화비유사성"], dl))
lg = M["S_유사성빈도"].copy()
np.fill_diagonal(lg, 0)
check("legacy diagonal-0 sheet = S with only the diagonal zeroed (a SIMILARITY count)", np.array_equal(M["기존_빈도행렬_대각0(유사성)"], lg))
check("legacy sheet differs from D off the diagonal (so it is not a dissimilarity matrix)", not np.array_equal(M["기존_빈도행렬_대각0(유사성)"], M["D_비유사성빈도"]))
wsP = wb_f["P_유사성비율"]
check("display format of P is 3 decimals but the stored value keeps full precision", wsP.cell(2, 3).number_format == "0.000" and any(M["P_유사성비율"][i, j] != round(M["P_유사성비율"][i, j], 3) for i in range(n) for j in range(n)))
if a.scope == "ALL":
    c = bundle["comparison"]
    check("S_KR + S_JP = S_ALL", np.array_equal(matrix(wb_v["S_한국"]) + matrix(wb_v["S_일본"]), M["S_유사성빈도"]) and np.array_equal(matrix(wb_v["S_한국"]), np.array(c["Skr"], dtype=float)))

# formulas: only on the check sheet, with cached values
formula_cells = {ws.title: [c.coordinate for r in ws.iter_rows() for c in r if c.data_type == "f"] for ws in wb_f.worksheets}
check("formulas exist only on 행렬정의·수식검산", all((not v) for k, v in formula_cells.items() if k != "행렬정의·수식검산") and len(formula_cells["행렬정의·수식검산"]) >= 4, {k: len(v) for k, v in formula_cells.items() if v})
wsv = wb_v["행렬정의·수식검산"]
cached = [(c.coordinate, c.value) for r in wsv.iter_rows() for c in r if c.coordinate in formula_cells["행렬정의·수식검산"]]
check("every formula has a cached result value", all(v is not None for _, v in cached), cached)
verdicts = [v for _, v in cached if isinstance(v, str)]
check("all formula verdicts say 일치", verdicts and all(v == "일치" for v in verdicts), verdicts)
nums = [v for _, v in cached if isinstance(v, (int, float))]
check("check sums are ~0 and the count equals n*n", sum(1 for v in nums if abs(v) < 1e-9) >= 3 and any(v == n * n for v in nums), nums)

# error values / literal formula text / hyperlinks
errs, injected, links = [], [], 0
for ws in wb_f.worksheets:
    links += len(ws._hyperlinks)
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str) and re.match(r"^#(REF|VALUE|DIV/0|NAME|N/A|NUM|NULL)", c.value):
                errs.append((ws.title, c.coordinate))
check("no error values", not errs, errs[:5])
check("no hyperlinks", links == 0, links)

# ids stay text
resp = wb_v["참여자_분류응답(가명)"]
rows = [[c.value for c in r] for r in resp.iter_rows(min_row=3)]  # row 1 = note, row 2 = header
hdr = [c.value for c in resp[2]]
check("response sheet header", hdr[:7] == ["가명ID", "국가", "제출시각(KST)", "진술문", "문항번호", "묶음번호(참여자 내)", "묶음 이름(참여자 입력)"], hdr)
resp_csv = read_csv("02_원자료_가명/참여자별_분류응답표.csv")
body = [r for r in rows if r[0] is not None]
check("response rows equal the CSV (count)", len(body) == len(resp_csv) - 1, (len(body), len(resp_csv) - 1))
ok_eq = True
for i, r in enumerate(body):
    cs = resp_csv[i + 1]
    if [str(r[0]), str(r[1]), str(r[2]), str(r[3]), str(int(r[4])), str(int(r[5])), "" if r[6] is None else str(r[6])] != cs:
        ok_eq = False
        break
check("response rows equal the CSV (every cell, same order)", ok_eq)
countries = sorted({r[1] for r in body})
want = ["KR"] if a.scope == "KR" else ["JP"] if a.scope == "JP" else ["JP", "KR"]
check("response sheet holds exactly the scope's country/countries", countries == want, countries)
per = {}
for r in body:
    per[r[0]] = per.get(r[0], 0) + 1
check(f"{N} participants x {n} statements", len(per) == N and all(v == n for v in per.values()), (len(per), set(per.values())))
check("participant ids and statement ids are text", all(isinstance(r[0], str) and isinstance(r[3], str) for r in body))
stm = wb_v["진술문_원문대응"]
shdr = [c.value for c in stm[1]]
if a.scope == "KR":
    check("KR statement sheet has Korean text only (no Japanese column)", not any("日本語" in str(h) or "일본어" in str(h) for h in shdr), shdr)
elif a.scope == "JP":
    check("JP statement sheet has the Japanese text first", "日本語" in str(shdr[3]), shdr)
else:
    check("ALL statement sheet has both languages", any("한국어" in str(h) for h in shdr) and any("日本語" in str(h) for h in shdr), shdr)
st_rows = [[c.value for c in r] for r in stm.iter_rows(min_row=2)]
check("47 statements in order S01..S47", [r[0] for r in st_rows] == [f"S{i + 1:02d}" for i in range(n)])
stm_csv = read_csv("03_진술문/진술문_대응표.csv")[1:]  # 번호, 순서, DB id, 한국어, 日本語, 검수
ko_ok = all(r[3 if a.scope != "JP" else 5] == c[3] for r, c in zip(st_rows, stm_csv))
check("Korean statement text equals the statement table CSV character for character", ko_ok)
if a.scope in ("JP", "ALL"):
    ja_col = 3 if a.scope == "JP" else 4
    check("Japanese statement text equals the statement table CSV character for character (kana/kanji preserved)", all(r[ja_col] == c[4] for r, c in zip(st_rows, stm_csv)) and any(re.search(r"[぀-ヿ]", str(r[ja_col])) for r in st_rows))

# coordinates / clusters
cs = wb_v["좌표·군집(추천)"]
chdr = [c.value for c in cs[2]]
crow = [[c.value for c in r] for r in cs.iter_rows(min_row=3)]
coords_csv = read_csv("05_MDS/좌표_2D.csv")
check("coordinates equal 좌표_2D.csv exactly (the CSV keeps full precision) and the bundle", all(float(coords_csv[i + 1][1]) == crow[i][2] and float(coords_csv[i + 1][2]) == crow[i][3] for i in range(n)) and np.allclose(np.array([[r[2], r[3]] for r in crow]), np.array(bundle["analysis"]["primary"]["coordinates"]), atol=1e-12, rtol=0))
rec = bundle["naming"]["proposal"]
sig = [int(x) for x in rec["solution"]["signature"].split(",")]
check("cluster membership equals the stored recommended solution", [r[4] for r in crow] == sig)
rc = wb_v["추천군집·문항"]
rrows = [[c.value for c in r] for r in rc.iter_rows(min_row=3)]
listed = [x.strip() for r in rrows if r[0] is not None for x in str(r[4]).split(",")]
check("each statement appears in exactly one recommended cluster", sorted(listed) == [f"S{i + 1:02d}" for i in range(n)], len(listed))
check("cluster sizes equal member counts", all(int(r[3]) == len(str(r[4]).split(",")) for r in rrows if r[0] is not None))
check("cluster names equal the stored proposal", {int(r[0]): r[1] for r in rrows if r[0] is not None} == {c["cluster"]: c["nameKo"] for c in rec["solution"]["clusters"]})
check("review status states 자동 제안·연구자 검토 전", all("연구자 검토 전" in str(r[9]) for r in rrows if r[0] is not None))

# pairs / dims / k table
pw = wb_v["문항쌍_진단"]
prow = [[c.value for c in r] for r in pw.iter_rows(min_row=3)]
pairs_csv = read_csv("06_진단/쌍진단_1081.csv")
check("1,081 pair rows", len(prow) == 1081 == len(pairs_csv) - 1, len(prow))
okp = all(prow[i][0] == pairs_csv[i + 1][0] and prow[i][1] == pairs_csv[i + 1][1] and all(abs(float(pairs_csv[i + 1][j]) - float(prow[i][j])) < 1e-12 for j in range(2, 11)) for i in range(len(prow)))
check("pair diagnostics equal 쌍진단_1081.csv", okp)
pr = bundle["analysis"]["pairDiagnostics"]
check("pair diagnostics equal the bundle", all(abs(pr[i]["distance"] - prow[i][6]) < 1e-12 and abs(pr[i]["residualKruskal"] - prow[i][10]) < 1e-12 for i in range(len(pr))))
dw = wb_v["차원별_지표"]
drow = [[c.value for c in r] for r in dw.iter_rows(min_row=2)]
dims_csv = read_csv("05_MDS/차원비교_1-5D.csv")
check("dimension table equals CSV and bundle", len(drow) == len(dims_csv) - 1 and all(abs(float(dims_csv[i + 1][1]) - drow[i][1]) < 1e-12 and abs(bundle["analysis"]["dimensions"][i]["stress1"] - drow[i][1]) < 1e-12 and abs(bundle["analysis"]["dimensions"][i]["rsq"] - drow[i][4]) < 1e-12 for i in range(len(drow))))
kw = wb_v["군집수별_지표"]
krow = [[c.value for c in r] for r in kw.iter_rows(min_row=2)]
k_csv = read_csv("07_군집/군집후보비교.csv")
check("k table equals CSV and bundle", len(krow) == len(k_csv) - 1 and all(abs(float(k_csv[i + 1][3]) - krow[i][3]) < 1e-9 and abs(bundle["analysis"]["solutions"][i]["silhouetteMap"]["mean"] - krow[i][3]) < 1e-5 for i in range(len(krow))))
sv = wb_v["분할반분_전체값"]
svrow = [[c.value for c in r] for r in sv.iter_rows(min_row=2)]
check("split-half values equal the bundle", len(svrow) == bundle["analysis"]["splitHalf"]["repetitions"] and all(abs(svrow[i][1] - bundle["analysis"]["splitHalf"]["rawValues"][i]) < 1e-12 for i in range(len(svrow)) if bundle["analysis"]["splitHalf"]["rawValues"][i] is not None))

# no personal-data columns / values
all_headers = " ".join(str(c.value) for ws in wb_v.worksheets for r in ws.iter_rows(min_row=1, max_row=3) for c in r if isinstance(c.value, str))
check("no name / phone / session-id / token / demographic columns", not re.search(r"참여자 ?이름|전화|phone|세션 ?ID|sessionId|adminToken|토큰|성별|학교급|학년\b|나이", all_headers.replace("가명", "")), "")

# charts native, no external links
z = zipfile.ZipFile(path)
parts = z.namelist()
check("native charts present", any(p.startswith("xl/charts/chart") for p in parts), [p for p in parts if "chart" in p][:4])
check("no external links / images / OLE", not any(p.startswith(("xl/externalLinks", "xl/media", "xl/embeddings")) for p in parts), [p for p in parts if p.startswith(("xl/externalLinks", "xl/media", "xl/embeddings"))])
rels = " ".join(z.read(p).decode("utf-8", "ignore") for p in parts if p.endswith(".rels"))
check("no http(s) targets in relationships", "http://" not in rels.replace("http://schemas.openxmlformats.org", "").replace("http://schemas.microsoft.com", "") and "https://" not in rels, "")

ok = all(r["pass"] for r in results)
json.dump({"workbook": os.path.basename(path), "scope": a.scope, "ok": ok, "checks": results, "sheets": names}, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"{sum(1 for r in results if r['pass'])}/{len(results)} checks passed — {os.path.basename(path)}")
for r in results:
    if not r["pass"]:
        print("FAIL", r["check"], r["detail"])
sys.exit(0 if ok else 1)
