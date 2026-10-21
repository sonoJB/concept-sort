"""Analysis report (Markdown + editable DOCX + PDF) composed from a stored bundle. Block list -> three renderers."""
import html
import os

import numpy as np

from common import SCOPE_LABEL_KO, fmt, ko_design, sid
from scopedata import ScopeData

CLAIM_SOURCE_ROWS = [
    ["개인 이진 유사성 행렬(대각선 1)과 집단 유사성 빈도 S(대각선 N), MDS→Ward 절차", "trochim1989", "Step 3, Step 4 (저자 공개 본문, 절 제목)", "전문 열람. 비유사성으로의 변환식(δ = 1 − S/N)은 이 본문에 없으며 이 분석이 명시한 변환식임(독립 재계산으로 검증)"],
    ["군집 수는 수학적 기준이 아닌 재량·내용 판단, 군집명은 합의로 제안, 지도 방향의 임의성과 군집 모양의 무의미", "trochim1989; trochim1994", "Trochim(1989) Step 4–5; Trochim 외(1994) Results·Session 2", "전문 열람. 두 논문 모두 신뢰도·Stress의 허용 기준을 제시하지 않음"],
    ["서열 SMACOF, 명시적 정규화 Σd̂² = n(n−1)/2, raw stress / Stress-1 식", "mair2022", "p. 4, 식 (2)–(4)", "전문 열람"],
    ["동점 처리 primary/secondary/tertiary", "mair2022", "pp. 3–4", "전문 열람"],
    ["임의 시작 여러 번과 국소 최소 문제", "mair2022", "p. 6, §3.1", "전문 열람"],
    ["Stress 경험칙(Kruskal 구술 척도)의 인용과 적용상 주의, 고정 기준값만으로 판단하지 말 것", "mair2016; mair2022; kruskal1964(2차 인용)", "Mair et al.(2016) 원고 pp. 3, 8–9; Mair et al.(2022) pp. 7–9", "전문 열람(Kruskal 원문은 미열람)"],
    ["동점이 많을수록·오차가 적을수록 서열 MDS Stress가 낮아짐", "mair2016", "원고 p. 9", "전문 열람"],
    ["선행 개념도 연구 69편의 Stress 분포: 평균 0.28, 표준편차 0.04, 관찰 범위 0.17–0.34 (비교 참고, 합격 기준 아님)", "rosas2012 (원문 미대조); rocliffe2023 (2차 문헌, 전문 열람)", "Rosas & Kane(2012) §3.5에 귀속; Rocliffe 외(2023) Results > Sorting/cluster map", "원문 접근 불가로 대조하지 못함. 평균·범위는 2차 문헌이 귀속, 표준편차 0.04는 연구자 제공값. 0.205–0.365(평균 0.285)는 별개의 비교 범위(일부 2차 문헌은 Kane & Trochim, 2007이 인용한 Trochim의 미발표 1993년 검토로 설명; Henderikx & Stoffers, 2023은 Rosas & Kane 인용으로 제시)이며 서로 다른 자료·범위일 수 있어 모순으로 단정하지 않음 → 출처 귀속은 원문 확인 필요"],
    ["RSQ 정의(디스패리티와 거리의 제곱 상관)", "multbiplotr (공식 문서)", "MDS 도움말 Value절", "양호 기준은 확인되지 않음"],
    ["Ward 최소분산 군집", "ward1963 (원문 미열람)", "위치 미확인", "독립 구현(군집 내 제곱합 증가량 직접 계산)과의 일치로 검증"],
    ["실루엣", "rousseeuw1987 (초록·서지만 확인)", "위치 미확인", "정의는 독립 구현과의 일치로 검증; 해석 구간은 확인하지 못함"],
    ["무작위 이분 + Spearman–Brown 보정 분할 신뢰도의 개념(사례 값 행렬 .79, 배치 .56)", "trochim1994", "Results > Reliability and Consistency of Maps", "전문 열람. 이 분석의 설계(1,000회, 국가 층화, seed 고정)는 사전 고정 설정이며 논문 설계가 아님; 허용 기준 없음"],
    ["순열 기반 Stress 규준(참고)", "mair2016", "원고 pp. 14–15", "이번 분석에서는 계산하지 않음(미실시)"],
]


def blocks_for(d: ScopeData, figs, tables, verif, figtext=None):
    """figs: {key: relative png path}; tables: ScopeData tables dict; verif: verification dict; figtext: {key: caption + note} of the figure list."""
    figtext = figtext or {}

    def cap(key, default):
        return figtext.get(key, default)

    a, m, scope = d.a, d.meta, d.scope
    p = a["primary"]
    sel = p["selected"]
    label = SCOPE_LABEL_KO[scope]
    B = []
    B.append(("h1", f"{label} 개념도 분석 보고서"))
    B.append(("p", f"분석 버전 {m['algorithmVersion']} · 소스 커밋 {m['engineSourceCommitSha'][:12]} · 실행 ID {m['runId']} · 자료 지문 {m['fingerprints']['participantsHash'][:16]}…"))
    B.append(("note", "이 보고서의 모든 수치는 저장된 분석 결과에서 읽거나 그 결과로부터 결정적으로 재계산한 값입니다. 학술지 게재 가능성이나 ‘매우 적합’ 같은 판정을 보장하지 않으며, 불리한 결과도 그대로 보고합니다. 군집 수·군집명은 문항 내용을 읽고 만든 자동 제안이며 연구자·참여자 합의를 거친 확정 명칭이 아닙니다."))

    # ---------------------------------------------------------------------------------------- 1
    B.append(("h2", "1. 자료"))
    prow = d.participant_rows()
    piles = [r["piles"] for r in prow]
    maxs = [r["max"] for r in prow]
    most = [r["pid"] for r in prow if r["piles"] == max(piles)]
    fewest = [r["pid"] for r in prow if r["piles"] == min(piles)]
    biggest = [r["pid"] for r in prow if r["max"] == max(maxs)]
    B.append(("ul", [
        f"대상: 본조사(dataRole=MAIN)만. 유효 참여자 N = {d.N} (한국 {m['nKr']}, 일본 {m['nJp']}), 진술문 {d.n}개. 파일럿 자료는 포함하지 않음.",
        f"자료 기준 시점: {m['dataSnapshotAt']} (UTC). 세 범위(한국·일본·전체)에 같은 기준 시점과 같은 포함·제외 규칙을 적용함.",
        "포함 규칙: 본조사, 국가 코드가 KR 또는 JP, 47개 진술문을 정확히 한 번씩 분류한 완결 제출. 제외 규칙: 파일럿·비본조사, 국가 코드 없음, 미완료, 진술문 중복 배치, 무효 진술문.",
        f"참여자별 묶음 수: 최소 {min(piles)}, 중앙값 {float(np.median(piles)):g}, 최대 {max(piles)}. 가장 큰 묶음의 크기: 최소 {min(maxs)}, 중앙값 {float(np.median(maxs)):g}, 최대 {max(maxs)}.",
        f"특이 응답: 임의로 수정·제외하지 않았음. 묶음 수가 가장 많은 참여자 {', '.join(most)}({max(piles)}개), 가장 적은 참여자 {', '.join(fewest)}({min(piles)}개), 가장 큰 묶음({max(maxs)}문항)을 만든 참여자 {len(biggest)}명. 전체 요약은 ‘참여자별묶음요약’ 표에 있음.",
        "번역: 일본어 대응문의 검수 상태는 시스템상 값만 확인함. 번역 동등성(예: 역번역) 검토 기록은 이 분석 자료에 없으므로 수행했다고 주장하지 않음.",
        "평정 자료: 중요도·실행가능성 평정은 수집되지 않았으므로 평정지도·패턴매칭·Go-Zone을 만들지 않았고 분류 빈도로 대체하지 않았음.",
    ]))
    B.append(("h3", "참여자 가명화"))
    B.append(("p", "개별 응답은 가명 연구 ID(KR-001 등, 제출 순서)로만 제공하며 이름·전화번호·실제 세션 ID·접근 토큰은 포함하지 않습니다. 분류 응답값(묶음 구성)과 참여자가 붙인 묶음 이름은 원본과 동일하게 보존했습니다."))

    # ---------------------------------------------------------------------------------------- 2
    B.append(("h2", "2. 방법"))
    B.append(("table", ["행렬", "정의", "대각선", "큰 값의 의미"], [["집단 유사성 빈도 S", "같은 묶음에 분류한 참여자 수 (개인 이진 유사성 행렬의 합)", "N", "더 유사함"], ["유사성 비율 P", "S / N", "1", "더 유사함"], ["비유사성 빈도 D", "N − S", "0", "더 다름"], ["비유사성 비율 δ", "1 − S / N", "0", "더 다름 — MDS 입력"]]))
    B.append(("p", "위 식은 분석 대상 참여자 전원이 모든 진술문을 완전히 분류했음을 확인한 뒤에만 적용했습니다. 결측을 0으로 바꾸거나 분모를 바꾸지 않았습니다. 유사성이 높은 문항일수록 δ가 작아 더 가깝게 배치되며, 입력 방향과 이중 변환이 없음을 계산으로 확인했습니다(δ = 1 − S/N을 독립적으로 다시 계산해 일치)."))
    s = a["settings"]["mds"]
    B.append(("ul", [
        f"MDS: 2차원 비계량(서열) SMACOF. 동점 처리 secondary(같은 δ는 같은 디스패리티), 가중치 모두 1(대각선 제외), 디스패리티 정규화 Σ d̂² = n(n−1)/2, 목적함수 raw stress = Σ(d̂ − d)². δ = 0인 쌍을 결측으로 처리하지 않음(실제 0 값 유지).",
        f"수렴 정의: (σ_(k−1) − σ_k)/σ_(k−1) ≤ ε. 분모는 직전 반복의 raw stress. raw stress의 증가(상대 1e-9 초과)는 OBJECTIVE_INCREASED, 반복 상한 도달은 MAX_ITER로 구분하며 Stress-1이 일시적으로 올라도 raw stress가 줄면 실패로 보지 않음.",
        f"계산 설정(이번 작업의 설정이며 문헌상 합격선이 아님): 재현 가능한 시작 {s['nInit']}회(seed {s['randomSeed']}에서 파생), 시작당 최대 {s['maxIter']}회, 허용오차 {s['eps']}. 오류 없이 수렴한 시작 중 raw stress가 가장 작은 해를 선택하고, 그 해에서 허용오차 {s['strictEps']}로 계속 최적화해 안정성을 확인함.",
        "여러 시작 중 최저 해를 고른 것은 전역 최적의 증명이 아님(Mair et al., 2022).",
        "최종 좌표에서 raw stress, Stress-1, RSQ를 다시 계산해 저장값과 일치함을 확인했고, 독립 구현(numpy)으로 교차 검증함.",
        "1~5차원은 같은 방법으로 계산한 보조 Stress 비교이며 3차원 이상 결과를 2차원 지도로 쓰지 않음.",
        f"군집분석: 2차원 지도 좌표에 대한 Ward 최소분산 계층 군집(Ward, 1963; Lance–Williams 갱신). 군집 수 k = {a['settings']['ward']['kMin']}–{a['settings']['ward']['kMax']}를 비교하고 지도 실루엣·δ 실루엣(Rousseeuw, 1987), 병합 높이·비용, 최소 군집 크기, 문항 내용을 함께 검토함.",
        f"반복 분할 신뢰도: 참여자를 무작위로 반분({ko_design(a['splitHalf']['design'])})해 문항 쌍(1,081개) 상관을 {a['splitHalf']['repetitions']}회(seed {a['splitHalf']['seed']}) 계산하고 Spearman–Brown 값을 함께 보고함. 행렬 신뢰도이며 지도 신뢰도가 아님.",
    ]))

    # ---------------------------------------------------------------------------------------- 3
    B.append(("h2", "3. 결과"))
    B.append(("h3", "3.1 입력 행렬 검증"))
    c = a["matrixChecks"]
    B.append(("p", f"행렬 크기 {c['size']}×{c['size']}, 대칭 {'예' if c['symmetric'] else '아니오'}, 대각선 = N {'예' if c['diagonalEqualsN'] else '아니오'}, 값 범위 정상 {'예' if c['offDiagonalInRange'] else '아니오'}. 비대각 최소 {c['minOffDiagonal']}, 최대 {c['maxOffDiagonal']}. 한 번도 함께 분류되지 않은 쌍 {c['neverTogether']}/{c['pairs']}, 모두가 함께 분류한 쌍 {c['alwaysTogether']}, 서로 다른 비유사성 수준 {c['distinctLevels']}개(동점이 많음). 참여자 기록에서 S를 다시 세어 저장값과 일치했고({verif['S_equals_recount_from_participant_records']}), 기존 Excel 행렬과 비대각 값이 일치함."))
    B.append(("h3", "3.2 MDS 수렴과 적합도 (둘은 별개)"))
    B.append(("table", ["차원", "Stress-1", "RSQ", "raw stress", "수렴 시작", "종료 사유(주/엄격)", "수치적 수렴"], [[f"{x['dimension']}D" + (" (지도)" if x["isPrimary"] else " (보조)"), fmt(x["stress1"], 4), fmt(x["rsq"], 3), fmt(x["rawStress"], 3), f"{x['convergedStarts']}/{x['startCount']}", f"{x['selected']['main']['reason']}/{x['selected']['strict']['reason']}" if x["selected"] else "", "예" if x["numericallyConverged"] else "아니오"] for x in a["dimensions"]]))
    B.append(("p", f"2차원 해: 주 실행 {sel['main']['reason']}({sel['main']['iterations']}회, 최종 변화량 {sel['main']['finalRelativeChange']:.2e}, ε={s['eps']}), 엄격 계속 {sel['strict']['reason']}({sel['strict']['iterations']}회, 최종 변화량 {sel['strict']['finalRelativeChange']:.2e}, ε={s['strictEps']}). 선택한 시작 #{sel['initIndex']}. {p['convergedStarts']}/{p['startCount']}개 시작이 수렴했고 최저 해의 1e-6 이내에 도달한 시작은 {p['startsWithinRelTol1e6OfBest']}개."))
    s1 = p["stress1"]
    fi = next((x for x in d.b["metrics"] if x["id"] == "fit_interpretation"), None)
    B.append(("h3", "2차원 표현의 적합도 해석 (3.2 계속)"))
    if fi:
        B.append(("p", fi["definitionKo"]))
        B.append(("ul", ["수렴(알고리즘이 국소 최소에 도달했는가): " + next(v["textKo"] for v in fi["verdicts"] if v["aspectKo"] == "수렴"), "2차원 표현의 적합도(Stress-1·raw stress·RSQ·잔차): " + next(v["textKo"] for v in fi["verdicts"] if v["aspectKo"] == "2차원 표현의 적합도"), "분류 안정성(반복 분할 신뢰도): 3.5절과 지표표 참조 — 양호 기준 미확정", "내용 타당성(군집·명칭과 문항 내용의 대응): 이 보고서가 판정하지 않으며 연구자 검토 대상"]))
    B.append(("p", f"Stress-1 = {s1:.4f}, RSQ = {p['rsq']:.3f}. Stress > 0.20이라는 사실만으로 개념도 분석을 실패·부적합으로 분류하지 않습니다(Kruskal 구술 척도의 기계적 적용 문제: Mair et al., 2022, §3.2). RSQ는 같은 정의의 양호 기준을 원문에서 확인하지 못해 ‘참고 지표·양호 기준 미확정’으로 둡니다. 3차원 Stress-1은 {fmt(next(x for x in a['dimensions'] if x['dimension'] == 3)['stress1'], 3)}이며, 차원이 늘면 Stress가 줄어드는 것은 일반 원칙이므로 이 감소만으로 2차원 표현의 부족을 결론짓지 않습니다."))
    if a["tieSensitivity"]:
        t = a["tieSensitivity"]
        B.append(("p", f"동점 처리 민감도(공식 결과 아님): primary(동점 분리 허용)로 같은 설정에서 계산하면 Stress-1 = {fmt(t['stress1'], 4)}, RSQ = {fmt(t['rsq'], 3)}. 동점이 많은 자료에서 동점 처리 선택이 Stress에 영향을 줍니다(Mair et al., 2016)."))
    for key, default in (("shepard", "Shepard 다이어그램"), ("resid", "잔차 진단"), ("itemstress", "문항별 stress 기여"), ("iter", "목적함수(raw stress) 반복 이력"), ("dims", "차원별 Stress-1·RSQ (보조 비교)")):
        if key in figs:
            B.append(("fig", figs[key], cap(key, default)))
    B.append(("h3", "3.3 2차원 개념도와 군집"))
    if d.sol:
        pr = d.proposal
        B.append(("p", f"추천 군집 수 {d.rec_k}(검토 상태: 자동 제안·연구자 검토 전). ‘최적’은 검토한 후보(k={a['settings']['ward']['kMin']}–{a['settings']['ward']['kMax']})와 아래 기준에서의 추천을 뜻하며 수학적 유일 정답이 아닙니다."))
        B.append(("ul", pr["recommendationRationaleKo"]))
        B.append(("fig", figs["map_plain"], cap("map_plain", "2차원 점 지도")))
        B.append(("fig", figs["map_color"], cap("map_color", f"개념도(2차원, {d.rec_k}군집, 색)") + " 지도의 방향은 임의이고 군집의 크기·모양에는 의미 있는 해석이 없다(Trochim et al., 1994)."))
        B.append(("fig", figs["map_gray"], cap("map_gray", "같은 지도의 흑백판(기호·해치로 군집 구분)")))
        bd = d.boundary(d.sol)
        rows = []
        for cl in d.sol["clusters"]:
            nm = d.names.get(cl["cluster"], {})
            bnd = ", ".join(f"S{i + 1}" for i in cl["members"] if bd[i]["is_boundary"])
            rows.append([cl["cluster"], nm.get("nameKo", ""), nm.get("nameJa", ""), cl["size"], ", ".join(f"S{i + 1}" for i in cl["members"]), bnd or "—"])
        B.append(("table", ["군집", "추천 군집명", "日本語(검토용 번역안)", "문항 수", "포함 문항", "경계 문항(지도 실루엣<0.1)"], rows))
        for cl in d.sol["clusters"]:
            nm = d.names.get(cl["cluster"], {})
            B.append(("p", f"군집 {cl['cluster']} · {nm.get('nameKo', '')} — 핵심 내용: {nm.get('coreContentKo', '')} 명명 근거: {nm.get('basisKo', '')}"))
        B.append(("p", "군집 소속·번호·명칭은 화면, Excel, 그림, 보고서에서 같은 값을 사용합니다. 군집명은 실제 참여자·전문가 합의를 거친 최종 명칭이 아닙니다. 번호는 군집의 가장 작은 문항 번호 순으로 부여해 지도의 방향(임의적)과 무관합니다."))
        overlaps = [(cl["cluster"], cl["foreignPointsInside"]) for cl in d.sol["clusters"] if cl["foreignPointsInside"]]
        B.append(("p", "경계 겹침: " + ("; ".join(f"군집 {c}의 경계 안에 문항 {', '.join(f'S{i + 1}' for i in v)}" for c, v in overlaps) if overlaps else "추천 해에서는 다른 군집의 점이 경계(볼록 껍질) 안에 들어온 경우가 없음.")))
        B.append(("h3", "3.4 군집 수 비교"))
        merges = {mm["clustersAfter"]: mm["height"] for mm in a["ward"]["merges"]}
        alt = {x["k"]: x["commentKo"] for x in pr["alternatives"]}
        B.append(("table", ["k", "최소 군집", "2문항 이하", "지도 실루엣", "δ 실루엣", "군집 간 제곱합 비율", "k+1→k 병합 높이", "비고"], [[sl["k"] if sl["k"] != d.rec_k else f"{sl['k']} ★", sl["minSize"], sl["smallClusters"], fmt(sl["silhouetteMap"]["mean"], 3), fmt(sl["silhouetteDelta"]["mean"], 3), fmt(sl["betweenShare"], 3), fmt(merges.get(sl["k"]), 3), alt.get(sl["k"], "")] for sl in a["solutions"]]))
        for key, default in (("kcmp", "군집 수 비교"), ("dendro", "Ward 덴드로그램"), ("sil", "문항별 실루엣(지도 좌표 기준)")):
            if key in figs:
                B.append(("fig", figs[key], cap(key, default)))
    else:
        B.append(("p", "군집 추천이 첨부되지 않았습니다."))
    B.append(("h3", "3.5 반복 분할 신뢰도"))
    sh = a["splitHalf"]
    B.append(("p", f"{ko_design(sh['design'])}, {sh['repetitions']}회(seed {sh['seed']}). 원상관 평균 {fmt(sh['rawR']['mean'], 3)}(중앙값 {fmt(sh['rawR']['median'], 3)}, 최소–최대 {fmt(sh['rawR']['min'], 3)}–{fmt(sh['rawR']['max'], 3)}), Spearman–Brown 평균 {fmt(sh['spearmanBrown']['mean'], 3)}. 계산 불능 {sh['undefinedCount']}회, 음수 {sh['negativeRawCount']}회. 분위수 범위는 이 표본의 무작위 분할 분포이며 모집단 신뢰구간이 아니고, 행렬 신뢰도를 지도 신뢰도로 바꾸어 해석하지 않습니다."))
    if "split" in figs:
        B.append(("fig", figs["split"], cap("split", "반복 분할 신뢰도 분포")))
    B.append(("p", "브리징 값: 산식과 정규화 방식을 학술지 원문에서 확인하지 못해 산출하지 않았습니다(임의 대체식을 공식 브리징 값이라고 부르지 않음)."))
    cmp_ = d.b.get("comparison")
    if cmp_:
        B.append(("h3", "3.6 한국·일본 비교(탐색적)"))
        B.append(("p", f"N이 다른 빈도를 직접 비교하지 않고 P_KR = S_KR/N_KR (N={cmp_['nKr']}), P_JP = S_JP/N_JP (N={cmp_['nJp']})를 사용했습니다. 동일 문항 쌍의 Pearson r = {cmp_['pearson']:.3f}, Spearman ρ = {cmp_['spearman']:.3f}, |P_KR−P_JP| 평균 {cmp_['meanAbsDiff']:.3f}, 최대 {cmp_['maxAbsDiff']:.3f}. 1,081개 문항 쌍은 독립 관측치가 아니므로 일반 상관검정의 p값은 제시하지 않으며, 전체는 한국·일본을 포함하므로 세 범위를 독립 표본처럼 비교하지 않습니다."))
        for key, default in (("cmp_heat", "P_KR − P_JP 히트맵(공통 색 범위)"), ("cmp_scatter", "P_KR 대 P_JP"), ("cmp_proc", "한국·일본 지도의 Procrustes 정렬(탐색적)")):
            if key in figs:
                B.append(("fig", figs[key], cap(key, default)))

    # ---------------------------------------------------------------------------------------- 4
    B.append(("h2", "4. 지표별 정의·기준·판정·출처·한계"))
    rows = []
    for me in d.b["metrics"]:
        crit = " / ".join(f"[{c['labelKo']}] {c['text']}" + (f" (위치: {c['location']})" if c.get("location") else "") for c in me["criteria"])
        ver = " / ".join(f"{v['aspectKo']}: {v['textKo']}" for v in me["verdicts"])
        rows.append([me["nameKo"], me["formula"], me["valueDisplay"], crit, ver, ", ".join(me["sourceKeys"]), " / ".join(me["limitsKo"])])
    B.append(("table", ["지표", "계산식", "실제 값", "기준(종류별)", "판정", "출처", "적용 한계"], rows))
    # ---------------------------------------------------------------------------------------- 5
    B.append(("h2", "5. 해석의 한계"))
    B.append(("ul", [
        "수치적 수렴은 국소 최소에 도달했다는 사실이며 2차원 표현이 자료에 적합하다는 뜻이 아님. 적합도는 Stress-1·RSQ·잔차 진단·안정성·내용 해석을 함께 검토해야 하며 일반 MDS의 고정 임계값만으로 판정하지 않음.",
        "여러 시작 중 최저 해이며 전역 최적의 증명이 아님. 동점이 많아(서로 다른 비유사성 수준이 적음) 동점 처리 선택에 결과가 민감할 수 있음.",
        "지도의 방향(회전·반사)은 임의적이며 축에는 실질적 의미를 부여하지 않았음. 의미를 부여하려면 별도 내용 근거와 연구자 검토가 필요함.",
        "군집 수와 군집명은 문항 내용을 읽고 만든 자동 제안이며 참여자·전문가 합의, 표집·번역·내용 타당도 검토는 이 분석에 포함되지 않음. 방법론 근거가 있어도 이들이 자동으로 확보되지는 않음.",
        "실루엣·분할 신뢰도의 양호 기준은 원문에서 확인하지 못해 판정하지 않았음. RSQ의 양호 기준도 미확정.",
        "Rosas & Kane(2012)의 Stress 분포(평균 0.28, 표준편차 0.04, 관찰 범위 0.17–0.34)는 원문을 직접 대조하지 못했고(2차 문헌 귀속, 표준편차는 연구자 제공값), 선행 연구의 Stress 산식·정규화·동점 처리가 본 분석의 Stress-1과 같은지도 확인되지 않음. 관찰 범위는 비교 참고이며 보편적 합격 기준이 아님.",
        "표본 크기(한국 30, 일본 38)는 작고, 한국·일본의 지도 비교는 정렬 절차 없이 해석할 수 없음.",
    ]))
    B.append(("h2", "6. 참고문헌 (APA 7판)과 열람 상태"))
    B.append(("refs", d.b["references"]))
    B.append(("h3", "주장–출처 대응표"))
    B.append(("table", ["주장·방법", "출처", "확인한 원문 위치", "비고"], CLAIM_SOURCE_ROWS))
    B.append(("h2", "7. 재현"))
    B.append(("p", "01_요약의 분석자료 Excel에는 이 보고서의 표(S·P·D·δ, 좌표·군집, 차원별·군집 수별 지표, 문항 쌍 진단, 안정성, 지표 정의·참고문헌)가 편집 가능한 실제 셀로 들어 있고, 08_그림의 그림목록.csv가 각 그림의 캡션·주석과 대응 데이터 파일을 연결합니다. 그림은 인쇄 폭 16 cm 기준 PNG 600 dpi와 SVG·PDF 벡터로 제공합니다(보고서 안의 그림은 파일 크기 때문에 절반 해상도 사본)."))
    B.append(("p", "10_재현 폴더에 설정(settings.json), 환경, 실행 방법, 독립 구현(reproduce_mds.py)과 TypeScript 소스를 포함했습니다. 02_원자료_가명의 참여자별 분류응답표만으로 S 행렬·MDS·Ward·분할 신뢰도를 다시 계산해 결과와 대조할 수 있습니다. 11_검증에 검증 보고서와 SHA-256 목록이 있습니다."))
    return number_figs(B)


def number_figs(B):
    import re

    out, n = [], 0
    for b in B:
        if b[0] == "fig":
            n += 1
            out.append(("fig", b[1], f"그림 {n}. " + re.sub(r"^그림 \d+\. ", "", b[2])))
        else:
            out.append(b)
    return out


# ------------------------------------------------------------------------------------------- renderers
def render_md(blocks, outdir_figs_prefix=""):
    L = []
    for b in blocks:
        t = b[0]
        if t == "h1":
            L.append(f"# {b[1]}\n")
        elif t == "h2":
            L.append(f"\n## {b[1]}\n")
        elif t == "h3":
            L.append(f"\n### {b[1]}\n")
        elif t == "p":
            L.append(b[1] + "\n")
        elif t == "note":
            L.append("> " + b[1] + "\n")
        elif t == "ul":
            L.extend(f"- {x}" for x in b[1])
            L.append("")
        elif t == "table":
            L.append("| " + " | ".join(b[1]) + " |")
            L.append("|" + "---|" * len(b[1]))
            for r in b[2]:
                L.append("| " + " | ".join(str(x).replace("|", "\\|").replace("\n", " ") for x in r) + " |")
            L.append("")
        elif t == "fig":
            L.append(f"![{b[2]}]({outdir_figs_prefix}{b[1]})\n\n*{b[2]}*\n")
        elif t == "refs":
            for r in b[1]:
                L.append(f"- {r['apa']}  \n  열람 상태: {r['access']} — {r['accessNoteKo']}")
            L.append("")
    return "\n".join(L)


CSS = """
@font-face { font-family: 'MG'; src: url('malgun.ttf'); }
@font-face { font-family: 'MG'; font-weight: bold; src: url('malgunbd.ttf'); }
body { font-family: 'MG'; font-size: 9pt; line-height: 1.35; }
h1 { font-size: 17pt; margin: 0 0 6pt 0; } h2 { font-size: 13pt; margin: 14pt 0 4pt 0; border-bottom: 0.6pt solid #888; }
h3 { font-size: 10.5pt; margin: 9pt 0 3pt 0; } p, li { margin: 2pt 0; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0; } td, th { border: 0.5pt solid #777; padding: 2pt 3pt; font-size: 7.5pt; vertical-align: top; }
th { background-color: #e5e7eb; } .note { background-color: #fff7e6; border: 0.5pt solid #e0a030; padding: 4pt; }
.cap { font-size: 8pt; color: #333; margin-bottom: 8pt; } img { width: 100%; }
"""


def render_html(blocks):
    esc = html.escape
    H = []
    for b in blocks:
        t = b[0]
        if t in ("h1", "h2", "h3"):
            H.append(f"<{t}>{esc(b[1])}</{t}>")
        elif t == "p":
            H.append(f"<p>{esc(b[1])}</p>")
        elif t == "note":
            H.append(f"<p class='note'>{esc(b[1])}</p>")
        elif t == "ul":
            H.append("<ul>" + "".join(f"<li>{esc(str(x))}</li>" for x in b[1]) + "</ul>")
        elif t == "table":
            H.append("<table><tr>" + "".join(f"<th>{esc(str(h))}</th>" for h in b[1]) + "</tr>" + "".join("<tr>" + "".join(f"<td>{esc(str(x))}</td>" for x in r) + "</tr>" for r in b[2]) + "</table>")
        elif t == "fig":
            H.append(f"<p><img src='{esc(os.path.basename(b[1]))}'/></p><p class='cap'>{esc(b[2])}</p>")
        elif t == "refs":
            H.append("<ul>" + "".join(f"<li>{esc(r['apa'].replace('*', ''))}<br/>열람 상태: {esc(r['access'])} — {esc(r['accessNoteKo'])}</li>" for r in b[1]) + "</ul>")
    return "<html><body>" + "\n".join(H) + "</body></html>"


def _strings(o):
    if isinstance(o, str):
        yield o
    elif isinstance(o, dict):
        for v in o.values():
            yield from _strings(v)
    elif isinstance(o, (list, tuple)):
        for v in o:
            yield from _strings(v)


def _subset_font(src, dst, chars):
    """Writes a font file containing only `chars` (plus ASCII): the report PDF then embeds ~100 KB instead of a whole CJK font."""
    from fontTools import subset
    from fontTools.ttLib import TTFont

    font = TTFont(src, fontNumber=0, lazy=True)
    opts = subset.Options()
    opts.layout_features = ["*"]
    opts.notdef_outline = True
    opts.name_IDs = ["*"]
    sub = subset.Subsetter(opts)
    sub.populate(text="".join(sorted(set(chars) | set(chr(c) for c in range(32, 127)))))
    sub.subset(font)
    font.save(dst)


def _font_archive(chars):
    """CJK font files for the PDF (MuPDF Story), subsetted to the characters of the report. Windows: Malgun Gothic; Linux server
    image: Noto Sans CJK (apt fonts-noto-cjk); override with PACKAGE_FONT_REGULAR / PACKAGE_FONT_BOLD. The files are written under
    the names the CSS refers to."""
    import shutil
    import tempfile

    import pymupdf

    cands = [
        (os.environ.get("PACKAGE_FONT_REGULAR"), os.environ.get("PACKAGE_FONT_BOLD")),
        ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),
        ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ]
    for reg, bold in cands:
        if reg and bold and os.path.exists(reg) and os.path.exists(bold):
            d = tempfile.mkdtemp(prefix="pkgfonts_")
            try:
                _subset_font(reg, os.path.join(d, "malgun.ttf"), chars)
                _subset_font(bold, os.path.join(d, "malgunbd.ttf"), chars)
            except Exception:  # fall back to the whole font files (larger PDF, same look)
                shutil.copyfile(reg, os.path.join(d, "malgun.ttf"))
                shutil.copyfile(bold, os.path.join(d, "malgunbd.ttf"))
            return pymupdf.Archive(d), d
    raise SystemExit("no CJK font found for the PDF report (set PACKAGE_FONT_REGULAR / PACKAGE_FONT_BOLD)")


def render_pdf(blocks, figdir, out_pdf):
    import pymupdf

    arch, fontdir = _font_archive(set("".join(_strings(blocks))))
    arch.add(figdir)
    story = pymupdf.Story(html=render_html(blocks), user_css=CSS, archive=arch)
    w = pymupdf.DocumentWriter(out_pdf)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (42, 42, -42, -50)
    more = True
    while more:
        dev = w.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        w.end_page()
    w.close()
    doc = pymupdf.open(out_pdf)
    pages = doc.page_count
    doc.set_metadata({"title": os.path.basename(out_pdf), "creationDate": "", "modDate": "", "producer": "", "creator": "", "author": "", "subject": "", "keywords": "", "format": "PDF 1.7", "trapped": "", "encryption": None})
    data = doc.tobytes(garbage=4, deflate=True, no_new_id=True)
    doc.close()
    with open(out_pdf, "wb") as fh:
        fh.write(data)
    return pages


def render_docx(blocks, figdir, out_docx):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    for b in blocks:
        t = b[0]
        if t == "h1":
            doc.add_heading(b[1], 0)
        elif t == "h2":
            doc.add_heading(b[1], 1)
        elif t == "h3":
            doc.add_heading(b[1], 2)
        elif t in ("p", "note"):
            doc.add_paragraph(b[1])
        elif t == "ul":
            for x in b[1]:
                doc.add_paragraph(str(x), style="List Bullet")
        elif t == "table":
            tb = doc.add_table(rows=1, cols=len(b[1]))
            tb.style = "Table Grid"
            for j, h in enumerate(b[1]):
                tb.rows[0].cells[j].text = str(h)
            for r in b[2]:
                cells = tb.add_row().cells
                for j, x in enumerate(r):
                    cells[j].text = str(x)
                    for par in cells[j].paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(8)
        elif t == "fig":
            doc.add_picture(os.path.join(figdir, os.path.basename(b[1])), width=Cm(15.5))
            doc.add_paragraph(b[2])
        elif t == "refs":
            for r in b[1]:
                doc.add_paragraph(f"{r['apa'].replace('*', '')} [열람 상태: {r['access']} — {r['accessNoteKo']}]", style="List Bullet")
    from datetime import datetime

    cp = doc.core_properties
    cp.created = cp.modified = cp.last_printed = datetime(2026, 1, 1)
    cp.author = cp.last_modified_by = "concept-sort analysis package"
    cp.title = os.path.basename(out_docx)
    doc.save(out_docx)
    from common import rewrite_zip_deterministic

    rewrite_zip_deterministic(out_docx)
