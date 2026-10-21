"""Small checks for the workbook writer (run: python -X utf8 test_workbook.py)."""
import os
import re
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402

import workbook as W  # noqa: E402


def make(path):
    B = W.Book(path)
    B.table("resp", ["가명ID", "묶음 이름"], [["KR-001", "=1+1"], ["KR-002", "+SUM(1,2)"], ["KR-003", "http://example.com/x"], ["KR-004", "@home"], ["KR-005", "정상"]], text_cols={0, 1})
    ws = B.sheet("meta")
    B.put(ws, 0, 0, 0.3382352941176471, B.f_3)
    B.put(ws, 1, 0, 1 / 3, B.f_3)
    B.put(ws, 2, 0, 0.1 + 0.2, B.f_3)
    B.put(ws, 3, 0, "=2+2")
    B.close()


def main():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "t.xlsx")
        make(p)
        with zipfile.ZipFile(p) as x:
            sheets = "".join(x.read(n).decode("utf-8") for n in x.namelist() if n.startswith("xl/worksheets/sheet"))
        assert "<f>" not in sheets, "a participant-typed string became a formula"
        assert "hyperlink" not in sheets, "a participant-typed string became a hyperlink"
        wb = openpyxl.load_workbook(p)
        vals = [c.value for c in wb[W.SHEET["resp"]]["B"][1:]]
        assert vals == ["=1+1", "+SUM(1,2)", "http://example.com/x", "@home", "정상"], vals
        assert all(c.data_type == "s" for c in wb[W.SHEET["resp"]]["B"][1:]), "text cells must stay strings"
        m = wb[W.SHEET["meta"]]
        for i, v in enumerate([0.3382352941176471, 1 / 3, 0.1 + 0.2]):
            assert m.cell(i + 1, 1).value == v, (i, m.cell(i + 1, 1).value, v)  # full double precision, not 16 digits
            assert m.cell(i + 1, 1).number_format == "0.000"
        assert m.cell(4, 1).data_type == "s" and m.cell(4, 1).value == "=2+2"
        wb.close()
    print("workbook writer: all checks passed")


main()
